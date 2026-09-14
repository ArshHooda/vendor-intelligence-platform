from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any, Callable, Iterable
from urllib import error, parse, request

import psycopg
from openpyxl import load_workbook
from psycopg.types.json import Jsonb


DEFAULT_FILES = "Bills972.xlsx,4DMTVendorListingResults775.xlsx"
DEFAULT_BUCKET = "ap-source-files"
VALID_SOURCE_TYPES = frozenset({"bills", "vendors"})
INT32_MAX = 2_147_483_647
IMPORT_LOCK_KEY = 972_775


@dataclass(frozen=True)
class ColumnInfo:
    name: str
    data_type: str
    udt_name: str
    is_nullable: bool
    default: str | None
    is_identity: bool
    is_generated: bool
    maximum_length: int | None

    @property
    def receives_server_value(self) -> bool:
        return self.default is not None or self.is_identity or self.is_generated


@dataclass(frozen=True)
class ConstraintInfo:
    name: str
    kind: str
    validated: bool
    definition: str
    columns: tuple[str, ...]
    referenced_schema: str | None
    referenced_table: str | None
    referenced_columns: tuple[str, ...]


@dataclass(frozen=True)
class PolicyInfo:
    name: str
    command: str
    roles: tuple[str, ...]
    using_expression: str | None
    check_expression: str | None


@dataclass(frozen=True)
class TableInfo:
    schema: str
    name: str
    columns: dict[str, ColumnInfo]
    constraints: tuple[ConstraintInfo, ...]
    row_security: bool
    force_row_security: bool
    owner: str
    policies: tuple[PolicyInfo, ...]
    privileges: dict[str, bool]
    standalone_unique_indexes: tuple[str, ...]
    user_triggers: tuple[str, ...]
    forbidden_privileges: tuple[str, ...]

    @property
    def qualified_name(self) -> str:
        return f"{self.schema}.{self.name}"


@dataclass(frozen=True)
class SourceRow:
    excel_row_number: int
    cell_values: tuple[Any, ...]


@dataclass(frozen=True)
class SheetData:
    name: str
    header_row_number: int
    headers: tuple[str, ...]
    rows: tuple[SourceRow, ...]


@dataclass(frozen=True)
class WorkbookData:
    object_path: str
    source_type: str
    file_sha256: str
    sheets: tuple[SheetData, ...]

    @property
    def row_count(self) -> int:
        return sum(len(sheet.rows) for sheet in self.sheets)


class SchemaCompatibilityError(RuntimeError):
    pass


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Import private Excel workbooks into Supabase staging tables."
    )
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument(
        "--schema-only",
        action="store_true",
        help="Validate the live database contract without downloading files.",
    )
    mode.add_argument(
        "--preflight-only",
        action="store_true",
        help="Validate the database and workbooks without inserting rows.",
    )
    return parser.parse_args()


def require_env(name: str) -> str:
    value = os.getenv(name)
    if not value:
        raise RuntimeError(f"Missing environment variable: {name}")
    return value


def source_file_list() -> list[str]:
    files = [
        item.strip()
        for item in os.getenv("SOURCE_FILES", DEFAULT_FILES).split(",")
        if item.strip()
    ]
    if not files:
        raise RuntimeError("SOURCE_FILES did not contain any object paths.")
    if len(files) != len(set(files)):
        raise RuntimeError("SOURCE_FILES contains a duplicate object path.")
    for object_path in files:
        if "\x00" in object_path:
            raise RuntimeError("SOURCE_FILES contains a NUL character.")
    return files


def source_type_map() -> dict[str, str]:
    raw = os.getenv("SOURCE_TYPE_MAP")
    if not raw:
        return {}
    try:
        mapping = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise RuntimeError("SOURCE_TYPE_MAP must be a JSON object.") from exc
    if not isinstance(mapping, dict):
        raise RuntimeError("SOURCE_TYPE_MAP must be a JSON object.")

    result: dict[str, str] = {}
    for key, value in mapping.items():
        if not isinstance(key, str) or not isinstance(value, str):
            raise RuntimeError("SOURCE_TYPE_MAP keys and values must be strings.")
        normalized = value.strip().lower()
        if normalized not in VALID_SOURCE_TYPES:
            raise RuntimeError(
                f"Invalid SOURCE_TYPE_MAP value for {key!r}: {value!r}. "
                "Allowed values are bills and vendors."
            )
        result[key.casefold()] = normalized
    return result


def infer_source_type(object_path: str, explicit_map: dict[str, str]) -> str:
    basename = PurePosixPath(object_path.replace("\\", "/")).name
    for candidate in (object_path.casefold(), basename.casefold()):
        if candidate in explicit_map:
            return explicit_map[candidate]

    normalized_name = re.sub(r"[^a-z0-9]+", "", basename.casefold())
    if "bill" in normalized_name:
        return "bills"
    if "vendor" in normalized_name:
        return "vendors"
    raise RuntimeError(
        f"Cannot determine source type for {object_path!r}. Set SOURCE_TYPE_MAP to "
        'a JSON object such as {"file.xlsx":"bills"}.'
    )


def download_storage_object(
    base_url: str, key: str, bucket: str, object_path: str
) -> bytes:
    url = (
        f"{base_url.rstrip('/')}/storage/v1/object/"
        f"{parse.quote(bucket, safe='')}/{parse.quote(object_path, safe='/')}"
    )
    headers = {"apikey": key}
    if not key.startswith(("sb_secret_", "sb_publishable_")):
        headers["Authorization"] = f"Bearer {key}"
    storage_request = request.Request(url, headers=headers)
    try:
        with request.urlopen(storage_request, timeout=60) as response:
            return response.read()
    except error.HTTPError as exc:
        raise RuntimeError(f"Storage HTTP {exc.code}: {object_path}") from exc
    except error.URLError as exc:
        raise RuntimeError(f"Storage connection failed: {object_path}") from exc


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def json_safe_excel_value(value: Any, context: str) -> Any:
    if value is None or isinstance(value, (bool, int)):
        return value
    if isinstance(value, str):
        if "\x00" in value:
            raise RuntimeError(f"NUL character found at {context}.")
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise RuntimeError(f"Non-finite number found at {context}.")
        return value
    if isinstance(value, Decimal):
        if not value.is_finite():
            raise RuntimeError(f"Non-finite decimal found at {context}.")
        return int(value) if value == value.to_integral_value() else float(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, bytes):
        return value.hex()
    raise RuntimeError(
        f"Unsupported Excel value type {type(value).__name__} at {context}."
    )


def normalized_header(value: Any, position: int) -> str:
    if is_blank(value):
        return f"column_{position}"
    header = " ".join(str(value).strip().split())
    if "\x00" in header:
        raise RuntimeError(f"NUL character found in header column {position}.")
    return header or f"column_{position}"


def parse_workbook(data: bytes, object_path: str, source_type: str) -> WorkbookData:
    if data[:4] != b"PK\x03\x04":
        raise RuntimeError(f"Unexpected file format: {object_path}")
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise RuntimeError(f"Cannot read Excel workbook {object_path}: {exc}") from exc

    parsed_sheets: list[SheetData] = []
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            header_values: tuple[Any, ...] | None = None
            header_row_number = 0
            for row_number, row in enumerate(rows, start=1):
                if any(not is_blank(cell) for cell in row):
                    header_values = tuple(row)
                    header_row_number = row_number
                    break

            if header_values is None:
                print(f"WARN: Ignoring empty sheet {object_path}:{worksheet.title}")
                continue

            headers = tuple(
                normalized_header(value, position)
                for position, value in enumerate(header_values, start=1)
            )
            if not headers:
                raise RuntimeError(f"No columns found in {object_path}:{worksheet.title}.")

            parsed_rows: list[SourceRow] = []
            for row_number, row in enumerate(rows, start=header_row_number + 1):
                raw_values = tuple(row)
                if all(is_blank(value) for value in raw_values):
                    continue
                values = tuple(
                    json_safe_excel_value(
                        raw_values[index] if index < len(raw_values) else None,
                        f"{object_path}:{worksheet.title}!R{row_number}C{index + 1}",
                    )
                    for index in range(len(headers))
                )
                parsed_rows.append(SourceRow(row_number, values))

            parsed_sheets.append(
                SheetData(
                    name=worksheet.title,
                    header_row_number=header_row_number,
                    headers=headers,
                    rows=tuple(parsed_rows),
                )
            )
    finally:
        workbook.close()

    if not parsed_sheets:
        raise RuntimeError(f"Workbook has no importable sheets: {object_path}")
    return WorkbookData(
        object_path=object_path,
        source_type=source_type,
        file_sha256=hashlib.sha256(data).hexdigest(),
        sheets=tuple(parsed_sheets),
    )


def table_info(
    conn: psycopg.Connection[Any], schema_name: str, table_name: str
) -> TableInfo:
    column_rows = conn.execute(
        """
        select column_name, data_type, udt_name, is_nullable, column_default,
               is_identity, is_generated, character_maximum_length
        from information_schema.columns
        where table_schema = %s and table_name = %s
        order by ordinal_position
        """,
        (schema_name, table_name),
    ).fetchall()
    if not column_rows:
        raise SchemaCompatibilityError(
            f"Table not found or not visible: {schema_name}.{table_name}"
        )
    columns = {
        row[0]: ColumnInfo(
            name=row[0],
            data_type=row[1],
            udt_name=row[2],
            is_nullable=(row[3] == "YES"),
            default=row[4],
            is_identity=(row[5] == "YES"),
            is_generated=(row[6] != "NEVER"),
            maximum_length=row[7],
        )
        for row in column_rows
    }

    constraint_rows = conn.execute(
        """
        select c.conname, c.contype, c.convalidated,
               pg_get_constraintdef(c.oid, true),
               array(
                   select a.attname::text
                   from unnest(c.conkey) with ordinality as k(attnum, position)
                   join pg_attribute a
                     on a.attrelid = c.conrelid and a.attnum = k.attnum
                   order by k.position
               ),
               referenced_ns.nspname,
               referenced_table.relname,
               array(
                   select a.attname::text
                   from unnest(c.confkey) with ordinality as k(attnum, position)
                   join pg_attribute a
                     on a.attrelid = c.confrelid and a.attnum = k.attnum
                   order by k.position
               )
        from pg_constraint c
        join pg_class own_table on own_table.oid = c.conrelid
        join pg_namespace own_ns on own_ns.oid = own_table.relnamespace
        left join pg_class referenced_table on referenced_table.oid = c.confrelid
        left join pg_namespace referenced_ns
          on referenced_ns.oid = referenced_table.relnamespace
        where own_ns.nspname = %s and own_table.relname = %s
        order by c.conname
        """,
        (schema_name, table_name),
    ).fetchall()
    constraints = tuple(
        ConstraintInfo(
            name=row[0],
            kind=row[1],
            validated=bool(row[2]),
            definition=row[3],
            columns=tuple(row[4] or ()),
            referenced_schema=row[5],
            referenced_table=row[6],
            referenced_columns=tuple(row[7] or ()),
        )
        for row in constraint_rows
    )

    security_row = conn.execute(
        """
        select c.relrowsecurity, c.relforcerowsecurity, pg_get_userbyid(c.relowner)
        from pg_class c
        join pg_namespace n on n.oid = c.relnamespace
        where n.nspname = %s and c.relname = %s
        """,
        (schema_name, table_name),
    ).fetchone()
    policy_rows = conn.execute(
        """
        select policyname, cmd, roles, qual, with_check
        from pg_policies
        where schemaname = %s and tablename = %s
        order by policyname
        """,
        (schema_name, table_name),
    ).fetchall()
    policies = tuple(
        PolicyInfo(
            name=row[0],
            command=row[1],
            roles=tuple(row[2] or ()),
            using_expression=row[3],
            check_expression=row[4],
        )
        for row in policy_rows
    )

    qualified = f"{schema_name}.{table_name}"
    privileges = {
        privilege: bool(
            conn.execute(
                "select has_table_privilege(current_user, %s, %s)",
                (qualified, privilege),
            ).fetchone()[0]
        )
        for privilege in ("SELECT", "INSERT", "UPDATE", "DELETE")
    }
    standalone_unique_indexes = tuple(
        row[0]
        for row in conn.execute(
            """
            select pg_get_indexdef(i.indexrelid)
            from pg_index i
            join pg_class t on t.oid = i.indrelid
            join pg_namespace n on n.oid = t.relnamespace
            left join pg_constraint c on c.conindid = i.indexrelid
            where n.nspname = %s and t.relname = %s
              and i.indisunique and c.oid is null
            order by i.indexrelid::regclass::text
            """,
            (schema_name, table_name),
        ).fetchall()
    )
    user_triggers = tuple(
        row[0]
        for row in conn.execute(
            """
            select pg_get_triggerdef(g.oid, true)
            from pg_trigger g
            join pg_class t on t.oid = g.tgrelid
            join pg_namespace n on n.oid = t.relnamespace
            where n.nspname = %s and t.relname = %s and not g.tgisinternal
            order by g.tgname
            """,
            (schema_name, table_name),
        ).fetchall()
    )
    forbidden_privileges = tuple(
        f"{row[0]}:{row[1]}"
        for row in conn.execute(
            """
            with grants as (
                select
                    case
                        when acl.grantee = 0 then 'PUBLIC'
                        else pg_get_userbyid(acl.grantee)
                    end as grantee_name,
                    acl.privilege_type
                from pg_class c
                join pg_namespace n on n.oid = c.relnamespace
                cross join lateral aclexplode(
                    coalesce(c.relacl, acldefault('r', c.relowner))
                ) as acl
                where n.nspname = %s and c.relname = %s
            )
            select grantee_name, privilege_type
            from grants
            where grantee_name in ('PUBLIC', 'anon', 'authenticated')
              and privilege_type in ('SELECT', 'INSERT', 'UPDATE', 'DELETE')
            order by grantee_name, privilege_type
            """,
            (schema_name, table_name),
        ).fetchall()
    )
    return TableInfo(
        schema=schema_name,
        name=table_name,
        columns=columns,
        constraints=constraints,
        row_security=bool(security_row[0]) if security_row else False,
        force_row_security=bool(security_row[1]) if security_row else False,
        owner=security_row[2] if security_row else "",
        policies=policies,
        privileges=privileges,
        standalone_unique_indexes=standalone_unique_indexes,
        user_triggers=user_triggers,
        forbidden_privileges=forbidden_privileges,
    )


def compact_sql(value: str | None) -> str:
    return re.sub(r"\s+", " ", value or "").strip().lower()


def quoted_literals(definition: str) -> set[str]:
    return {
        match.replace("''", "'")
        for match in re.findall(r"'((?:''|[^'])*)'", definition)
    }


def constraint_with_columns(
    table: TableInfo, kind: str, columns: tuple[str, ...]
) -> ConstraintInfo | None:
    return next(
        (
            constraint
            for constraint in table.constraints
            if constraint.kind == kind and constraint.columns == columns
        ),
        None,
    )


def validate_checks(
    table: TableInfo,
    matchers: tuple[tuple[str, Callable[[str], bool]], ...],
    errors: list[str],
) -> None:
    checks = [constraint for constraint in table.constraints if constraint.kind == "c"]
    matched_names: set[str] = set()
    for label, matcher in matchers:
        matches = [
            constraint
            for constraint in checks
            if matcher(compact_sql(constraint.definition))
        ]
        if len(matches) != 1:
            errors.append(
                f"{table.qualified_name} must have exactly one {label} CHECK; "
                f"found {len(matches)}."
            )
        else:
            matched_names.add(matches[0].name)
            if not matches[0].validated:
                errors.append(
                    f"{table.qualified_name}.{matches[0].name} is not validated."
                )
    unexpected = [check.name for check in checks if check.name not in matched_names]
    if unexpected:
        errors.append(
            f"{table.qualified_name} has unsupported extra CHECK constraints: "
            + ", ".join(unexpected)
        )


def validate_database_contract(
    conn: psycopg.Connection[Any], tables: dict[str, TableInfo]
) -> None:
    errors: list[str] = []
    expected_columns: dict[str, dict[str, str]] = {
        "quality.pipeline_runs": {
            "run_id": "uuid",
            "started_at": "timestamp with time zone",
            "finished_at": "timestamp with time zone",
            "status": "text",
            "error_message": "text",
        },
        "staging.source_files": {
            "file_id": "uuid",
            "run_id": "uuid",
            "source_type": "text",
            "bucket_name": "text",
            "object_path": "text",
            "file_sha256": "text",
            "sheet_name": "text",
            "headers": "jsonb",
            "row_count": "integer",
            "imported_at": "timestamp with time zone",
        },
        "staging.source_rows": {
            "file_id": "uuid",
            "excel_row_number": "integer",
            "cell_values": "jsonb",
        },
    }
    inserted_columns = {
        "quality.pipeline_runs": set(),
        "staging.source_files": {
            "run_id",
            "source_type",
            "bucket_name",
            "object_path",
            "file_sha256",
            "sheet_name",
            "headers",
            "row_count",
        },
        "staging.source_rows": {"file_id", "excel_row_number", "cell_values"},
    }
    required_not_null = {
        "quality.pipeline_runs": {"run_id", "started_at", "status"},
        "staging.source_files": {
            "file_id", "run_id", "source_type", "bucket_name", "object_path",
            "file_sha256", "sheet_name", "headers", "row_count", "imported_at",
        },
        "staging.source_rows": {"file_id", "excel_row_number", "cell_values"},
    }
    required_defaults = {
        "quality.pipeline_runs": {"run_id", "started_at", "status"},
        "staging.source_files": {"file_id", "bucket_name", "imported_at"},
        "staging.source_rows": set(),
    }
    required_default_fragments = {
        "quality.pipeline_runs": {
            "run_id": "gen_random_uuid()",
            "started_at": "now()",
            "status": "running",
        },
        "staging.source_files": {
            "file_id": "gen_random_uuid()",
            "bucket_name": "ap-source-files",
            "imported_at": "now()",
        },
        "staging.source_rows": {},
    }
    required_privileges = {
        "quality.pipeline_runs": {"SELECT", "INSERT", "UPDATE"},
        "staging.source_files": {"SELECT", "INSERT"},
        "staging.source_rows": {"SELECT", "INSERT"},
    }

    for qualified_name, expected in expected_columns.items():
        table = tables[qualified_name]
        for name, expected_type in expected.items():
            column = table.columns.get(name)
            if column is None:
                errors.append(f"{qualified_name} is missing column {name}.")
                continue
            if column.data_type != expected_type:
                errors.append(
                    f"{qualified_name}.{name} must be {expected_type}; "
                    f"found {column.data_type}."
                )
            if name in required_not_null[qualified_name] and column.is_nullable:
                errors.append(f"{qualified_name}.{name} must be NOT NULL.")
            if name in required_defaults[qualified_name] and not column.receives_server_value:
                errors.append(f"{qualified_name}.{name} must have a server default.")
            expected_fragment = required_default_fragments[qualified_name].get(name)
            if expected_fragment and expected_fragment not in compact_sql(column.default):
                errors.append(
                    f"{qualified_name}.{name} default must contain "
                    f"{expected_fragment!r}; found {column.default!r}."
                )

        for column in table.columns.values():
            if (
                not column.is_nullable
                and not column.receives_server_value
                and column.name not in inserted_columns[qualified_name]
            ):
                errors.append(
                    f"{qualified_name}.{column.name} is NOT NULL with no default and "
                    "has no importer mapping."
                )
        actual_privileges = {
            privilege for privilege, enabled in table.privileges.items() if enabled
        }
        if actual_privileges != required_privileges[qualified_name]:
            errors.append(
                f"Current role privileges on {qualified_name} must be "
                f"{sorted(required_privileges[qualified_name])}; found "
                f"{sorted(actual_privileges)}."
            )
        if table.standalone_unique_indexes:
            errors.append(
                f"{qualified_name} has unsupported standalone unique indexes: "
                + " | ".join(table.standalone_unique_indexes)
            )
        if table.user_triggers:
            errors.append(
                f"{qualified_name} has unsupported user triggers: "
                + " | ".join(table.user_triggers)
            )
        if table.forbidden_privileges:
            errors.append(
                f"{qualified_name} exposes forbidden table privileges: "
                + ", ".join(table.forbidden_privileges)
            )

    for schema_name in ("quality", "staging"):
        allowed = conn.execute(
            "select has_schema_privilege(current_user, %s, 'USAGE')",
            (schema_name,),
        ).fetchone()[0]
        if not allowed:
            errors.append(f"Current role lacks USAGE on schema {schema_name}.")

    pipeline = tables["quality.pipeline_runs"]
    source_files = tables["staging.source_files"]
    source_rows = tables["staging.source_rows"]
    expected_relational = {
        "quality.pipeline_runs": {
            ("p", ("run_id",), None, None, ()),
        },
        "staging.source_files": {
            ("p", ("file_id",), None, None, ()),
            ("u", ("run_id", "source_type", "sheet_name"), None, None, ()),
            ("f", ("run_id",), "quality", "pipeline_runs", ("run_id",)),
        },
        "staging.source_rows": {
            ("p", ("file_id", "excel_row_number"), None, None, ()),
            ("f", ("file_id",), "staging", "source_files", ("file_id",)),
        },
    }
    for qualified_name, expected in expected_relational.items():
        table = tables[qualified_name]
        actual = {
            (
                constraint.kind,
                constraint.columns,
                constraint.referenced_schema,
                constraint.referenced_table,
                constraint.referenced_columns,
            )
            for constraint in table.constraints
            if constraint.kind in {"p", "u", "f", "x"}
        }
        for missing in expected - actual:
            errors.append(f"{qualified_name} is missing required constraint {missing!r}.")
        for extra in actual - expected:
            errors.append(f"{qualified_name} has unsupported constraint {extra!r}.")

    validate_checks(
        pipeline,
        (
            (
                "status allow-list",
                lambda d: "status" in d
                and quoted_literals(d) == {"running", "succeeded", "failed"},
            ),
        ),
        errors,
    )
    validate_checks(
        source_files,
        (
            (
                "source_type allow-list",
                lambda d: "source_type" in d
                and quoted_literals(d) == {"bills", "vendors"},
            ),
            (
                "SHA-256 format",
                lambda d: "file_sha256" in d and "^[0-9a-f]{64}$" in d,
            ),
            ("headers JSON array", lambda d: "jsonb_typeof(headers)" in d and "array" in d),
            ("non-negative row_count", lambda d: "row_count" in d and ">= 0" in d),
        ),
        errors,
    )
    validate_checks(
        source_rows,
        (
            ("Excel row lower bound", lambda d: "excel_row_number" in d and ">= 2" in d),
            ("cell_values JSON array", lambda d: "jsonb_typeof(cell_values)" in d and "array" in d),
        ),
        errors,
    )

    current_user = conn.execute("select current_user").fetchone()[0].casefold()
    role_settings = conn.execute(
        """
        select rolsuper, rolcreatedb, rolcreaterole, rolinherit,
               rolreplication, rolbypassrls
        from pg_roles
        where rolname = current_user
        """
    ).fetchone()
    if role_settings is None:
        errors.append("Current database role is not visible in pg_roles.")
    else:
        setting_names = (
            "SUPERUSER", "CREATEDB", "CREATEROLE", "INHERIT", "REPLICATION", "BYPASSRLS"
        )
        enabled_settings = [
            name for name, enabled in zip(setting_names, role_settings) if enabled
        ]
        if enabled_settings:
            errors.append(
                "Importer role has disallowed role attributes: "
                + ", ".join(enabled_settings)
            )
    for table in tables.values():
        if table.owner.casefold() == current_user:
            errors.append(
                f"Importer role must not own {table.qualified_name}; owners can bypass RLS."
            )
        if not table.row_security:
            errors.append(f"{table.qualified_name} must have row-level security enabled.")
            continue
        applicable = [
            policy
            for policy in table.policies
            if policy.command == "ALL"
            and current_user in {role.casefold() for role in policy.roles}
            and compact_sql(policy.using_expression).strip("()") == "true"
            and compact_sql(policy.check_expression).strip("()") == "true"
        ]
        if not applicable:
            errors.append(
                f"{table.qualified_name} needs an ALL RLS policy for the current role "
                "with USING (true) and WITH CHECK (true)."
            )

    if errors:
        details = "\n".join(
            f"  {index}. {message}" for index, message in enumerate(errors, 1)
        )
        raise SchemaCompatibilityError(
            "Database schema preflight found incompatible settings:\n" + details
        )


def schema_fingerprint(tables: dict[str, TableInfo]) -> str:
    payload = [
        {
            "table": name,
            "columns": [
                (column.name, column.data_type, column.is_nullable, column.default)
                for column in table.columns.values()
            ],
            "constraints": [
                (constraint.name, constraint.kind, constraint.definition)
                for constraint in table.constraints
            ],
            "rls": table.row_security,
            "policies": [
                (policy.name, policy.command, policy.roles)
                for policy in table.policies
            ],
        }
        for name, table in sorted(tables.items())
    ]
    encoded = json.dumps(payload, sort_keys=True, default=str).encode()
    return hashlib.sha256(encoded).hexdigest()[:12]


def load_database_contract(
    conn: psycopg.Connection[Any],
) -> dict[str, TableInfo]:
    tables = {
        "quality.pipeline_runs": table_info(conn, "quality", "pipeline_runs"),
        "staging.source_files": table_info(conn, "staging", "source_files"),
        "staging.source_rows": table_info(conn, "staging", "source_rows"),
    }
    current_user = conn.execute("select current_user").fetchone()[0]
    print(f"SCHEMA: role={current_user} fingerprint={schema_fingerprint(tables)}")
    for name, table in sorted(tables.items()):
        privileges = ",".join(k for k, v in table.privileges.items() if v) or "none"
        print(
            f"SCHEMA: {name} columns={len(table.columns)} "
            f"rls={str(table.row_security).lower()} privileges={privileges}"
        )
        for constraint in table.constraints:
            print(f"CONSTRAINT: {name}.{constraint.name} {constraint.definition}")
    validate_database_contract(conn, tables)
    probe_database_contract(conn, tables)
    print(
        "PREFLIGHT: database schema, constraints, RLS, privileges, and "
        "rollback probe are compatible"
    )
    return tables


def validate_import_plan(workbooks: Iterable[WorkbookData]) -> None:
    errors: list[str] = []
    sheet_keys: set[tuple[str, str]] = set()
    for workbook in workbooks:
        if workbook.source_type not in VALID_SOURCE_TYPES:
            errors.append(f"{workbook.object_path}: invalid source type.")
        if not re.fullmatch(r"[0-9a-f]{64}", workbook.file_sha256):
            errors.append(f"{workbook.object_path}: invalid SHA-256 fingerprint.")
        for sheet in workbook.sheets:
            key = (workbook.source_type, sheet.name)
            if key in sheet_keys:
                errors.append(f"Duplicate source_type/sheet_name in one run: {key!r}.")
            sheet_keys.add(key)
            if not sheet.name or "\x00" in sheet.name:
                errors.append(f"{workbook.object_path}: invalid worksheet name.")
            if not sheet.headers or any(not isinstance(h, str) or not h for h in sheet.headers):
                errors.append(f"{workbook.object_path}:{sheet.name}: invalid headers array.")
            if len(sheet.rows) > INT32_MAX:
                errors.append(f"{workbook.object_path}:{sheet.name}: row count exceeds integer range.")
            try:
                json.dumps(list(sheet.headers), allow_nan=False)
            except (TypeError, ValueError) as exc:
                errors.append(f"{workbook.object_path}:{sheet.name}: invalid headers JSON: {exc}")
            seen_rows: set[int] = set()
            for row in sheet.rows:
                if row.excel_row_number < 2 or row.excel_row_number > INT32_MAX:
                    errors.append(
                        f"{workbook.object_path}:{sheet.name}: invalid Excel row "
                        f"{row.excel_row_number}."
                    )
                if row.excel_row_number in seen_rows:
                    errors.append(
                        f"{workbook.object_path}:{sheet.name}: duplicate Excel row "
                        f"{row.excel_row_number}."
                    )
                seen_rows.add(row.excel_row_number)
                if len(row.cell_values) != len(sheet.headers):
                    errors.append(
                        f"{workbook.object_path}:{sheet.name}: row {row.excel_row_number} "
                        f"has {len(row.cell_values)} values for {len(sheet.headers)} headers."
                    )
                try:
                    json.dumps(list(row.cell_values), allow_nan=False)
                except (TypeError, ValueError) as exc:
                    errors.append(
                        f"{workbook.object_path}:{sheet.name}: row {row.excel_row_number} "
                        f"is invalid JSON: {exc}"
                    )
    if errors:
        details = "\n".join(
            f"  {index}. {message}" for index, message in enumerate(errors, 1)
        )
        raise RuntimeError("Workbook preflight failed:\n" + details)


def load_workbooks(
    base_url: str, key: str, bucket: str, object_paths: list[str]
) -> list[WorkbookData]:
    explicit_map = source_type_map()
    workbooks: list[WorkbookData] = []
    for object_path in object_paths:
        source_type = infer_source_type(object_path, explicit_map)
        workbook = parse_workbook(
            download_storage_object(base_url, key, bucket, object_path),
            object_path,
            source_type,
        )
        workbooks.append(workbook)
        print(
            f"PREFLIGHT: {object_path} source_type={source_type} "
            f"sheets={len(workbook.sheets)} rows={workbook.row_count} "
            f"sha256={workbook.file_sha256[:12]}..."
        )
        for sheet in workbook.sheets:
            print(
                f"PREFLIGHT:   sheet={sheet.name!r} header_row={sheet.header_row_number} "
                f"columns={len(sheet.headers)} rows={len(sheet.rows)}"
            )
    validate_import_plan(workbooks)
    print("PREFLIGHT: workbook structure and values are compatible")
    return workbooks


def expected_sheet_keys(workbooks: Iterable[WorkbookData]) -> set[tuple[str, str, str, str]]:
    return {
        (workbook.object_path, workbook.source_type, workbook.file_sha256, sheet.name)
        for workbook in workbooks
        for sheet in workbook.sheets
    }


def complete_prior_run(
    conn: psycopg.Connection[Any], bucket: str, workbooks: list[WorkbookData]
) -> Any | None:
    expected = expected_sheet_keys(workbooks)
    rows = conn.execute(
        """
        select f.run_id, f.object_path, f.source_type, f.file_sha256, f.sheet_name,
               f.row_count, count(rows.file_id) as stored_row_count
        from staging.source_files f
        join quality.pipeline_runs r on r.run_id = f.run_id
        left join staging.source_rows rows on rows.file_id = f.file_id
        where f.bucket_name = %s and r.status = 'succeeded'
        group by f.run_id, f.object_path, f.source_type, f.file_sha256,
                 f.sheet_name, f.row_count, r.finished_at
        order by r.finished_at desc nulls last
        """,
        (bucket,),
    ).fetchall()
    by_run: dict[Any, set[tuple[str, str, str, str]]] = {}
    for run_id, object_path, source_type, sha256, sheet_name, expected_rows, stored_rows in rows:
        if expected_rows != stored_rows:
            print(
                f"WARN: Ignoring incomplete prior sheet in run {run_id}: "
                f"{object_path}:{sheet_name} expected={expected_rows} stored={stored_rows}"
            )
            continue
        by_run.setdefault(run_id, set()).add(
            (object_path, source_type, sha256, sheet_name)
        )
    return next((run_id for run_id, keys in by_run.items() if expected <= keys), None)


def start_pipeline_run(conn: psycopg.Connection[Any]) -> Any:
    row = conn.execute(
        "insert into quality.pipeline_runs default values returning run_id"
    ).fetchone()
    if row is None:
        raise RuntimeError("Pipeline run insert returned no run_id.")
    return row[0]


def finish_pipeline_run(
    conn: psycopg.Connection[Any], run_id: Any, status: str, error_message: str | None
) -> None:
    result = conn.execute(
        """
        update quality.pipeline_runs
        set finished_at = %s, status = %s, error_message = %s
        where run_id = %s
        """,
        (datetime.now(timezone.utc), status, error_message, run_id),
    )
    if result.rowcount != 1:
        raise RuntimeError(f"Pipeline status update affected {result.rowcount} rows.")


def insert_source_file(
    conn: psycopg.Connection[Any],
    run_id: Any,
    bucket: str,
    workbook: WorkbookData,
    sheet: SheetData,
) -> Any:
    row = conn.execute(
        """
        insert into staging.source_files (
            run_id, source_type, bucket_name, object_path, file_sha256,
            sheet_name, headers, row_count
        ) values (%s, %s, %s, %s, %s, %s, %s, %s)
        returning file_id
        """,
        (
            run_id,
            workbook.source_type,
            bucket,
            workbook.object_path,
            workbook.file_sha256,
            sheet.name,
            Jsonb(list(sheet.headers)),
            len(sheet.rows),
        ),
    ).fetchone()
    if row is None:
        raise RuntimeError(
            f"Source file insert returned no file_id for {workbook.object_path}:{sheet.name}."
        )
    return row[0]


def insert_source_rows(
    conn: psycopg.Connection[Any], file_id: Any, sheet: SheetData, batch_size: int
) -> None:
    query = """
        insert into staging.source_rows (file_id, excel_row_number, cell_values)
        values (%s, %s, %s)
    """
    for start in range(0, len(sheet.rows), batch_size):
        batch = sheet.rows[start : start + batch_size]
        conn.executemany(
            query,
            [
                (file_id, row.excel_row_number, Jsonb(list(row.cell_values)))
                for row in batch
            ],
        )


def probe_database_contract(
    conn: psycopg.Connection[Any], tables: dict[str, TableInfo]
) -> None:
    probe_sheet = SheetData(
        name="__schema_probe__",
        header_row_number=1,
        headers=("probe",),
        rows=(SourceRow(2, (None,)),),
    )
    probe_workbook = WorkbookData(
        object_path="__schema_probe__.xlsx",
        source_type="bills",
        file_sha256="0" * 64,
        sheets=(probe_sheet,),
    )
    try:
        with conn.transaction(force_rollback=True):
            run_id = start_pipeline_run(conn)
            file_id = insert_source_file(
                conn, run_id, DEFAULT_BUCKET, probe_workbook, probe_sheet
            )
            insert_source_rows(conn, file_id, probe_sheet, 1)
            conn.execute("set constraints all immediate")
            finish_pipeline_run(conn, run_id, "succeeded", None)
    except psycopg.Error as exc:
        raise SchemaCompatibilityError(
            describe_database_error(exc, tables, "Rollback-only schema probe")
        ) from None


def describe_database_error(
    exc: psycopg.Error, tables: dict[str, TableInfo], context: str
) -> str:
    sqlstate = exc.sqlstate or "unknown"
    diag = exc.diag
    category = {
        "23502": "NOT NULL violation",
        "23503": "foreign-key violation",
        "23505": "unique/primary-key violation",
        "23514": "CHECK constraint violation",
        "22001": "value too long",
        "22P02": "invalid value representation",
        "22003": "numeric value out of range",
        "42501": "permission or row-level-security failure",
    }.get(sqlstate, "database error")
    parts = [f"{context}: {category} (SQLSTATE {sqlstate})"]
    if diag.schema_name or diag.table_name:
        parts.append(f"relation={diag.schema_name or '?'}.{diag.table_name or '?'}")
    if diag.column_name:
        parts.append(f"column={diag.column_name}")
    if diag.constraint_name:
        parts.append(f"constraint={diag.constraint_name}")
        matching = next(
            (
                constraint
                for table in tables.values()
                for constraint in table.constraints
                if constraint.name == diag.constraint_name
            ),
            None,
        )
        if matching:
            parts.append(f"definition={matching.definition}")
    if getattr(diag, "message_primary", None):
        parts.append(f"message={diag.message_primary}")
    return "; ".join(parts)


def import_workbooks(
    conn: psycopg.Connection[Any],
    tables: dict[str, TableInfo],
    workbooks: list[WorkbookData],
    bucket: str,
    batch_size: int,
) -> dict[str, int]:
    acquired = conn.execute(
        "select pg_try_advisory_lock(%s)", (IMPORT_LOCK_KEY,)
    ).fetchone()[0]
    if not acquired:
        raise RuntimeError("Another source-file import is already running.")
    try:
        return _import_workbooks_locked(conn, tables, workbooks, bucket, batch_size)
    finally:
        try:
            conn.execute("select pg_advisory_unlock(%s)", (IMPORT_LOCK_KEY,))
        except psycopg.Error as exc:
            print(f"WARN: Could not release advisory lock: {exc}", file=sys.stderr)


def _import_workbooks_locked(
    conn: psycopg.Connection[Any],
    tables: dict[str, TableInfo],
    workbooks: list[WorkbookData],
    bucket: str,
    batch_size: int,
) -> dict[str, int]:
    stale = conn.execute(
        """
        update quality.pipeline_runs
        set status = 'failed',
            finished_at = now(),
            error_message = 'Automatically marked failed after exceeding one hour'
        where status = 'running'
          and started_at < now() - interval '1 hour'
        """
    ).rowcount
    if stale:
        print(f"RECOVER: marked {stale} stale pipeline run(s) as failed")

    prior_run_id = complete_prior_run(conn, bucket, workbooks)
    if prior_run_id is not None:
        print(f"SKIP: identical complete import already succeeded as run {prior_run_id}")
        return {
            "files_seen": len(workbooks),
            "files_loaded": 0,
            "files_skipped": len(workbooks),
            "sheets_loaded": 0,
            "rows_loaded": 0,
        }

    metrics = {
        "files_seen": len(workbooks),
        "files_loaded": 0,
        "files_skipped": 0,
        "sheets_loaded": 0,
        "rows_loaded": 0,
    }
    try:
        run_id = start_pipeline_run(conn)
    except psycopg.Error as exc:
        raise RuntimeError(
            describe_database_error(exc, tables, "Create pipeline run")
        ) from None

    try:
        with conn.transaction():
            for workbook in workbooks:
                for sheet in workbook.sheets:
                    file_id = insert_source_file(conn, run_id, bucket, workbook, sheet)
                    insert_source_rows(conn, file_id, sheet, batch_size)
                metrics["files_loaded"] += 1
                metrics["sheets_loaded"] += len(workbook.sheets)
                metrics["rows_loaded"] += workbook.row_count
                print(
                    f"LOAD: {workbook.object_path} source_type={workbook.source_type} "
                    f"sheets={len(workbook.sheets)} rows={workbook.row_count}"
                )
            conn.execute("set constraints all immediate")
            finish_pipeline_run(conn, run_id, "succeeded", None)
        return metrics
    except Exception as exc:
        message = (
            describe_database_error(exc, tables, "Import transaction")
            if isinstance(exc, psycopg.Error)
            else str(exc)
        )
        try:
            finish_pipeline_run(conn, run_id, "failed", message[:10000])
        except Exception as update_exc:
            print(
                f"WARN: Could not mark pipeline run {run_id} as failed: {update_exc}",
                file=sys.stderr,
            )
        raise RuntimeError(message) from None


def main() -> int:
    args = parse_args()
    try:
        batch_size = int(os.getenv("IMPORT_BATCH_SIZE", "1000"))
    except ValueError as exc:
        raise RuntimeError("IMPORT_BATCH_SIZE must be an integer.") from exc
    if not 1 <= batch_size <= 10_000:
        raise RuntimeError("IMPORT_BATCH_SIZE must be between 1 and 10000.")

    with psycopg.connect(autocommit=True) as conn:
        tables = load_database_contract(conn)
        if args.schema_only:
            return 0

        base_url = require_env("SUPABASE_URL")
        key = require_env("SUPABASE_SECRET_KEY")
        bucket = os.getenv("SOURCE_BUCKET", DEFAULT_BUCKET).strip()
        if not bucket or "\x00" in bucket:
            raise RuntimeError("SOURCE_BUCKET is invalid.")
        workbooks = load_workbooks(base_url, key, bucket, source_file_list())
        if args.preflight_only:
            return 0

        metrics = import_workbooks(conn, tables, workbooks, bucket, batch_size)
        print(json.dumps(metrics, sort_keys=True))
        return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"FAIL: {exc}", file=sys.stderr)
        raise SystemExit(1)
