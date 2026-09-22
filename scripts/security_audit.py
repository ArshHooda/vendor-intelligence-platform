"""Run repeatable safety checks against source and generated dashboard data."""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

from dashboard_data.security import validate_public_output


UNSAFE_BROWSER_PATTERNS = {
    "HTML injection sink": re.compile(r"\b(?:innerHTML|outerHTML|insertAdjacentHTML)\b"),
    "dynamic code execution": re.compile(r"\b(?:eval|Function)\s*\("),
    "document.write": re.compile(r"\bdocument\.write\s*\("),
}
SECRET_PATTERNS = {
    "Supabase secret key": re.compile(r"\bsb_secret_[A-Za-z0-9_-]{16,}\b"),
    "private key": re.compile(r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----"),
    "database password URL": re.compile(r"postgres(?:ql)?://[^\s/:]+:[^\s@]+@", re.IGNORECASE),
}


def audit_frontend(root: Path) -> list[str]:
    failures: list[str] = []
    frontend_files = [root / "dist" / "index.html", root / "dist" / "app.js"]
    frontend_files.extend((root / "dist" / "modules").glob("*.js"))
    for path in frontend_files:
        content = path.read_text(encoding="utf-8")
        for label, pattern in UNSAFE_BROWSER_PATTERNS.items():
            if pattern.search(content):
                failures.append(f"{path.relative_to(root)} uses {label}")
    index = (root / "dist" / "index.html").read_text(encoding="utf-8")
    for required in ("Content-Security-Policy", 'type="module"', 'name="referrer"'):
        if required not in index:
            failures.append(f"dist/index.html is missing {required}")
    return failures


def audit_secrets(root: Path) -> list[str]:
    failures: list[str] = []
    paths = [
        path
        for path in root.rglob("*")
        if path.is_file()
        and path.resolve() != Path(__file__).resolve()
        and ".git" not in path.parts
        and "dist/data" not in path.as_posix()
        and path.suffix.lower() in {".py", ".js", ".html", ".css", ".md", ".sql", ".yml", ".yaml"}
    ]
    for path in paths:
        content = path.read_text(encoding="utf-8", errors="ignore")
        for label, pattern in SECRET_PATTERNS.items():
            if pattern.search(content):
                failures.append(f"{path.relative_to(root)} appears to contain a {label}")
    return failures


def public_string_values(value: Any) -> set[str]:
    if isinstance(value, str):
        return {value}
    if isinstance(value, dict):
        return set().union(*(public_string_values(item) for item in value.values()))
    if isinstance(value, list):
        return set().union(*(public_string_values(item) for item in value))
    return set()


def public_string_fields(value: Any, field: str = "root") -> dict[str, set[str]]:
    locations: dict[str, set[str]] = {}
    if isinstance(value, str):
        locations[value] = {field}
    elif isinstance(value, dict):
        for key, item in value.items():
            for text_value, fields in public_string_fields(item, key).items():
                locations.setdefault(text_value, set()).update(fields)
    elif isinstance(value, list):
        for item in value:
            for text_value, fields in public_string_fields(item, field).items():
                locations.setdefault(text_value, set()).update(fields)
    return locations


def audit_source_intersections(
    payload: dict[str, Any],
    workbook_path: Path,
    sensitive_headers: set[str],
    allowed_headers: set[str],
) -> list[str]:
    failures: list[str] = []
    public_values = public_string_values(payload.get("vendors", []))
    public_fields = public_string_fields(payload.get("vendors", []))
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = tuple(str(value or "").strip() for value in next(rows))
        sensitive_columns = [
            (index, header)
            for index, header in enumerate(headers)
            if header in sensitive_headers
        ]
        allowed_columns = [
            index for index, header in enumerate(headers) if header in allowed_headers
        ]
        source_values: dict[str, set[str]] = {
            header: set() for _, header in sensitive_columns
        }
        allowed_values: set[str] = set()
        for row in rows:
            for index in allowed_columns:
                value = row[index] if index < len(row) else None
                if isinstance(value, str) and len(value.strip()) >= 6:
                    allowed_values.add(value.strip())
                elif isinstance(value, (date, datetime)):
                    allowed_values.add(value.date().isoformat() if isinstance(value, datetime) else value.isoformat())
            for index, header in sensitive_columns:
                value = row[index] if index < len(row) else None
                if isinstance(value, str):
                    normalized = value.strip()
                    if len(normalized) >= 6:
                        source_values[header].add(normalized)
        for header, values in source_values.items():
            matches = (values - allowed_values) & public_values
            count = len(matches)
            if count:
                matched_fields = sorted(
                    set().union(*(public_fields[value] for value in matches))
                )
                failures.append(
                    f"Public output repeats {count} value(s) from sensitive source column "
                    f"{header} under public field(s): {', '.join(matched_fields)}."
                )
    finally:
        workbook.close()
    return failures


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", type=Path)
    parser.add_argument("--bills-file", type=Path)
    parser.add_argument("--vendors-file", type=Path)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    root = Path(__file__).resolve().parents[1]
    failures = audit_frontend(root) + audit_secrets(root)

    if args.data:
        payload = json.loads(args.data.read_text(encoding="utf-8"))
        try:
            validate_public_output(payload)
        except ValueError as exc:
            failures.append(str(exc))
        if args.bills_file:
            failures.extend(
                audit_source_intersections(
                    payload,
                    args.bills_file,
                    {
                        "Document Number",
                        "Transaction Number",
                        "Account",
                        "Memo",
                        "Preferred Entity Bank",
                        "Entity Bank (Vendor)",
                        "Entity Bank (Employee)",
                        "Entity Bank (Customer)",
                        "Entity Bank (Customer Credit)",
                        "Approver",
                        "Next Approver (EQ PR)",
                        "Therapeutic Area",
                        "Program",
                        "Indication",
                    },
                    {"Date", "Name", "Status", "Currency", "Payment Hold Reason"},
                )
            )
        if args.vendors_file:
            failures.extend(
                audit_source_intersections(
                    payload,
                    args.vendors_file,
                    {
                        "Address 1",
                        "Address 2",
                        "City",
                        "State/Province",
                        "Zip Code",
                        "Email Address",
                        "Internal ID",
                        "Comments",
                    },
                    {"ID", "Name", "Country", "Status", "Approval Status"},
                )
            )

    if failures:
        print("Security audit failed:", file=sys.stderr)
        for failure in failures:
            print(f"- {failure}", file=sys.stderr)
        return 1

    print("Security audit passed: public data allow-list, browser sinks, CSP, and secrets.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
