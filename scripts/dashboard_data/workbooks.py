"""Parse the source workbooks into safe intermediate records."""

from __future__ import annotations

from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path
from typing import Any, BinaryIO

import openpyxl

from .constants import CURRENCY_CODES, VENDOR_ID_PATTERN
from .values import (
    as_date,
    as_decimal,
    clean_text,
    header_positions,
    safe_purchase_reference,
    value_at,
)


def load_vendor_master(
    source: Path | BinaryIO,
) -> tuple[dict[str, dict[str, Any]], Counter]:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = header_positions(next(rows))
        vendors: dict[str, dict[str, Any]] = {}
        duplicates: Counter = Counter()

        for row in rows:
            vendor_id = clean_text(value_at(row, headers, "ID"))
            if not vendor_id:
                continue
            vendor_id = vendor_id.upper()
            duplicates[vendor_id] += 1
            last_paid = as_date(value_at(row, headers, "Last Paid Date"))
            candidate = {
                "vendor_source_id": vendor_id,
                "vendor_name": clean_text(value_at(row, headers, "Name")) or vendor_id,
                "vendor_status": clean_text(value_at(row, headers, "Status"))
                or "Not specified",
                "vendor_approval_status": clean_text(
                    value_at(row, headers, "Approval Status")
                )
                or "Not specified",
                "vendor_country": clean_text(value_at(row, headers, "Country"))
                or "Not specified",
                "master_last_paid_date": last_paid.isoformat() if last_paid else None,
            }
            existing = vendors.get(vendor_id)
            if not existing or sum(bool(value) for value in candidate.values()) > sum(
                bool(value) for value in existing.values()
            ):
                vendors[vendor_id] = candidate
        return vendors, duplicates
    finally:
        workbook.close()


def load_bill_activity(
    source: Path | BinaryIO,
    *,
    reference_salt: bytes,
) -> tuple[dict[str, dict[str, Any]], dict[str, int]]:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = header_positions(next(rows))
        activity: dict[str, dict[str, Any]] = defaultdict(
            lambda: {
                "bill_count": 0,
                "total_spend": Decimal("0"),
                "largest_bill_amount": None,
                "open_spend": Decimal("0"),
                "payment_hold_count": 0,
                "payment_hold_amount": Decimal("0"),
                "dates": [],
                "payment_status_counts": Counter(),
                "purchases": [],
            }
        )
        diagnostics = {
            "bill_rows": 0,
            "matched_bill_rows": 0,
            "unmatched_bill_rows": 0,
        }

        for row in rows:
            if not any(value is not None for value in row):
                continue
            diagnostics["bill_rows"] += 1
            bill_vendor = clean_text(value_at(row, headers, "Name")) or ""
            match = VENDOR_ID_PATTERN.search(bill_vendor)
            if not match:
                diagnostics["unmatched_bill_rows"] += 1
                continue

            diagnostics["matched_bill_rows"] += 1
            vendor_id = match.group(0).upper()
            base_amount = as_decimal(value_at(row, headers, "Amount"))
            original_amount = as_decimal(
                value_at(row, headers, "Amount (Foreign Currency)")
            )
            bill_date = as_date(value_at(row, headers, "Date"))
            payment_status = clean_text(value_at(row, headers, "Status")) or "Not specified"
            hold_reason = clean_text(value_at(row, headers, "Payment Hold Reason"))
            currency = clean_text(value_at(row, headers, "Currency")) or "Not specified"

            vendor = activity[vendor_id]
            vendor["bill_count"] += 1
            vendor["total_spend"] += base_amount
            vendor["largest_bill_amount"] = (
                base_amount
                if vendor["largest_bill_amount"] is None
                else max(vendor["largest_bill_amount"], base_amount)
            )
            vendor["payment_status_counts"][payment_status] += 1
            if "open" in payment_status.casefold():
                vendor["open_spend"] += base_amount
            if bill_date:
                vendor["dates"].append(bill_date)
            if hold_reason:
                vendor["payment_hold_count"] += 1
                vendor["payment_hold_amount"] += base_amount

            vendor["purchases"].append(
                {
                    "purchase_id": safe_purchase_reference(
                        value_at(row, headers, "Document Number"),
                        value_at(row, headers, "Transaction Number"),
                        salt=reference_salt,
                    ),
                    "purchase_date": bill_date.isoformat() if bill_date else None,
                    "status": payment_status,
                    "currency": currency,
                    "currency_code": CURRENCY_CODES.get(currency.casefold(), currency),
                    "original_amount": round(float(original_amount), 2),
                    "base_amount_usd": round(float(base_amount), 2),
                    "payment_hold": bool(hold_reason),
                    "payment_hold_reason": hold_reason,
                }
            )

        return dict(activity), diagnostics
    finally:
        workbook.close()
