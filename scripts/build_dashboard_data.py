"""Build a browser-safe vendor activity dataset for the static dashboard.

The script reads the two source workbooks either from local paths or from the
private Supabase Storage bucket. It publishes vendor-level aggregates only;
individual invoices, document numbers, memos, and bank fields are never written
to the dashboard artifact.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from statistics import mean
from typing import Any, BinaryIO

import openpyxl
from openpyxl.utils.datetime import from_excel


VENDOR_ID_PATTERN = re.compile(r"\bVEN\d{5}\b", re.IGNORECASE)
DEFAULT_BILLS_FILE = "Bills972.xlsx"
DEFAULT_VENDORS_FILE = "4DMTVendorListingResults775.xlsx"
DEFAULT_BUCKET = "ap-source-files"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--bills-file", type=Path)
    parser.add_argument("--vendors-file", type=Path)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("dist/data/vendor_activity.json"),
    )
    return parser.parse_args()


def storage_file(filename: str) -> BytesIO:
    base_url = os.getenv("SUPABASE_URL", "").rstrip("/")
    secret_key = os.getenv("SUPABASE_SECRET_KEY", "")
    bucket = os.getenv("SOURCE_BUCKET", DEFAULT_BUCKET)
    if not base_url or not secret_key:
        raise RuntimeError(
            "Set SUPABASE_URL and SUPABASE_SECRET_KEY, or pass both local workbook paths."
        )

    quoted_bucket = urllib.parse.quote(bucket, safe="")
    quoted_name = urllib.parse.quote(filename, safe="")
    request = urllib.request.Request(
        f"{base_url}/storage/v1/object/{quoted_bucket}/{quoted_name}",
        headers={"apikey": secret_key},
    )
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            content = response.read()
    except urllib.error.HTTPError as exc:
        raise RuntimeError(f"Storage returned HTTP {exc.code} for {filename}.") from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(f"Could not download {filename} from Supabase Storage.") from exc

    if content[:4] != b"PK\x03\x04":
        raise RuntimeError(f"Downloaded file is not an XLSX workbook: {filename}")
    return BytesIO(content)


def workbook_source(path: Path | None, filename: str) -> Path | BinaryIO:
    if path:
        if not path.is_file():
            raise FileNotFoundError(path)
        return path
    return storage_file(filename)


def header_positions(headers: tuple[Any, ...]) -> dict[str, int]:
    positions: dict[str, int] = {}
    for index, value in enumerate(headers):
        key = str(value or "").strip().casefold()
        if key and key not in positions:
            positions[key] = index
    return positions


def value_at(row: tuple[Any, ...], positions: dict[str, int], name: str) -> Any:
    index = positions.get(name.casefold())
    return row[index] if index is not None and index < len(row) else None


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.casefold() in {"- none -", "none", "null", "nan"}:
        return None
    return text


def as_decimal(value: Any) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, AttributeError):
        return Decimal("0")


def as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value)
            return converted.date() if isinstance(converted, datetime) else converted
        except (TypeError, ValueError, OverflowError):
            return None
    text = clean_text(value)
    if not text:
        return None
    for parser in (
        lambda item: datetime.fromisoformat(item.replace("Z", "+00:00")).date(),
        lambda item: datetime.strptime(item, "%m/%d/%Y").date(),
        lambda item: datetime.strptime(item, "%Y-%m-%d").date(),
    ):
        try:
            return parser(text)
        except ValueError:
            continue
    return None


def calendar_months(first_date: date, last_date: date) -> int:
    return max(
        1,
        (last_date.year - first_date.year) * 12
        + last_date.month
        - first_date.month
        + 1,
    )


def load_vendor_master(source: Path | BinaryIO) -> tuple[dict[str, dict[str, Any]], Counter]:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = header_positions(next(rows))
    vendors: dict[str, dict[str, Any]] = {}
    duplicates: Counter = Counter()

    for row in rows:
        vendor_id = clean_text(value_at(row, headers, "ID"))
        if not vendor_id:
            continue
        vendor_id = vendor_id.upper()
        duplicates[vendor_id] += 1
        candidate = {
            "vendor_source_id": vendor_id,
            "vendor_name": clean_text(value_at(row, headers, "Name")) or vendor_id,
            "vendor_status": clean_text(value_at(row, headers, "Status")) or "Not specified",
            "vendor_approval_status": clean_text(value_at(row, headers, "Approval Status"))
            or "Not specified",
            "vendor_country": clean_text(value_at(row, headers, "Country")) or "Not specified",
            "master_last_paid_date": (
                as_date(value_at(row, headers, "Last Paid Date")).isoformat()
                if as_date(value_at(row, headers, "Last Paid Date"))
                else None
            ),
        }
        existing = vendors.get(vendor_id)
        if not existing or sum(bool(value) for value in candidate.values()) > sum(
            bool(value) for value in existing.values()
        ):
            vendors[vendor_id] = candidate

    workbook.close()
    return vendors, duplicates


def load_bill_activity(source: Path | BinaryIO) -> tuple[dict[str, dict[str, Any]], dict[str, int]]:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = header_positions(next(rows))
    activity: dict[str, dict[str, Any]] = defaultdict(
        lambda: {
            "bill_count": 0,
            "total_spend": Decimal("0"),
            "largest_bill_amount": None,
            "open_spend": Decimal("0"),
            "payment_hold_count": 0,
            "payment_hold_amount": Decimal("0"),
            "dates": [],
            "payment_status_counts": Counter(),
        }
    )
    diagnostics = {"bill_rows": 0, "matched_bill_rows": 0, "unmatched_bill_rows": 0}

    for row in rows:
        if not any(value is not None for value in row):
            continue
        diagnostics["bill_rows"] += 1
        bill_vendor = clean_text(value_at(row, headers, "Name")) or ""
        match = VENDOR_ID_PATTERN.search(bill_vendor)
        if not match:
            diagnostics["unmatched_bill_rows"] += 1
            continue
        diagnostics["matched_bill_rows"] += 1
        vendor_id = match.group(0).upper()
        amount = as_decimal(value_at(row, headers, "Amount"))
        bill_date = as_date(value_at(row, headers, "Date"))
        payment_status = clean_text(value_at(row, headers, "Status")) or "Not specified"
        hold_reason = clean_text(value_at(row, headers, "Payment Hold Reason"))
        vendor = activity[vendor_id]
        vendor["bill_count"] += 1
        vendor["total_spend"] += amount
        vendor["largest_bill_amount"] = (
            amount
            if vendor["largest_bill_amount"] is None
            else max(vendor["largest_bill_amount"], amount)
        )
        vendor["payment_status_counts"][payment_status] += 1
        if "open" in payment_status.casefold():
            vendor["open_spend"] += amount
        if bill_date:
            vendor["dates"].append(bill_date)
        if hold_reason:
            vendor["payment_hold_count"] += 1
            vendor["payment_hold_amount"] += amount

    workbook.close()
    return dict(activity), diagnostics


def build_output(
    vendors: dict[str, dict[str, Any]],
    duplicates: Counter,
    activity: dict[str, dict[str, Any]],
    diagnostics: dict[str, int],
) -> dict[str, Any]:
    all_dates = [day for vendor in activity.values() for day in vendor["dates"]]
    as_of_date = max(all_dates) if all_dates else None
    total_spend = sum(
        (vendor["total_spend"] for vendor in activity.values()),
        start=Decimal("0"),
    )
    records: list[dict[str, Any]] = []

    for vendor_id, master in vendors.items():
        source = activity.get(vendor_id)
        bill_count = int(source["bill_count"]) if source else 0
        vendor_spend = source["total_spend"] if source else Decimal("0")
        dates = sorted(source["dates"]) if source else []
        first_purchase = dates[0] if dates else None
        last_purchase = dates[-1] if dates else None
        gaps = [(current - previous).days for previous, current in zip(dates, dates[1:])]
        status_counts = source["payment_status_counts"] if source else Counter()
        attention_status = next(
            (
                status
                for keyword in ("open", "reject")
                for status in status_counts
                if keyword in status.casefold()
            ),
            None,
        )
        primary_status = (
            attention_status
            or (
                sorted(status_counts.items(), key=lambda item: (-item[1], item[0]))[0][0]
                if status_counts
                else "No purchases"
            )
        )
        hold_count = int(source["payment_hold_count"]) if source else 0
        spend_share = float(vendor_spend / total_spend * 100) if total_spend else 0.0

        if duplicates[vendor_id] > 1:
            risk_label = "Duplicate vendor record"
        elif hold_count:
            risk_label = "Payment hold"
        elif spend_share >= 25:
            risk_label = "High concentration"
        elif not bill_count:
            risk_label = "No purchase history"
        else:
            risk_label = "Normal"

        records.append(
            {
                **master,
                "bill_count": bill_count,
                "total_spend": round(float(vendor_spend), 2),
                "average_bill_amount": round(float(vendor_spend / bill_count), 2)
                if bill_count
                else None,
                "largest_bill_amount": round(float(source["largest_bill_amount"]), 2)
                if source and source["largest_bill_amount"] is not None
                else None,
                "first_purchase_date": first_purchase.isoformat() if first_purchase else None,
                "last_purchase_date": last_purchase.isoformat() if last_purchase else None,
                "days_since_last_purchase": (as_of_date - last_purchase).days
                if as_of_date and last_purchase
                else None,
                "purchase_frequency_per_month": round(
                    bill_count / calendar_months(first_purchase, last_purchase), 2
                )
                if first_purchase and last_purchase
                else 0,
                "average_gap_days": round(mean(gaps), 1) if gaps else None,
                "payment_status": primary_status,
                "payment_statuses": sorted(status_counts),
                "open_spend": round(float(source["open_spend"]), 2) if source else 0,
                "payment_hold_count": hold_count,
                "payment_hold_amount": round(float(source["payment_hold_amount"]), 2)
                if source
                else 0,
                "vendor_record_count": duplicates[vendor_id],
                "spend_share_percent": round(spend_share, 4),
                "risk_label": risk_label,
            }
        )

    records.sort(key=lambda item: (item["vendor_name"].casefold(), item["vendor_source_id"]))
    return {
        "metadata": {
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "as_of_date": as_of_date.isoformat() if as_of_date else None,
            "vendor_master_rows": sum(duplicates.values()),
            "unique_vendors": len(vendors),
            "vendors_with_purchases": sum(record["bill_count"] > 0 for record in records),
            "total_spend": round(float(total_spend), 2),
            **diagnostics,
        },
        "vendors": records,
    }


def main() -> int:
    args = parse_args()
    try:
        vendors, duplicates = load_vendor_master(
            workbook_source(args.vendors_file, DEFAULT_VENDORS_FILE)
        )
        activity, diagnostics = load_bill_activity(
            workbook_source(args.bills_file, DEFAULT_BILLS_FILE)
        )
        output = build_output(vendors, duplicates, activity, diagnostics)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(
            json.dumps(output, ensure_ascii=False, separators=(",", ":")),
            encoding="utf-8",
        )
    except (FileNotFoundError, RuntimeError, ValueError, OSError) as exc:
        print(f"FAIL: {exc}", file=sys.stderr)
        return 1

    metadata = output["metadata"]
    print(
        "Dashboard data ready: "
        f"{metadata['unique_vendors']} vendors, "
        f"{metadata['matched_bill_rows']} matched bills, "
        f"${metadata['total_spend']:,.2f} total spend."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
